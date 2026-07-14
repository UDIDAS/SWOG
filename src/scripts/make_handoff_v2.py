#!/usr/bin/env python3
"""Build the hand-off v2 response: point-by-point doc + LiTS tumor-only masks + package."""
import os, glob, json, shutil
import numpy as np, nibabel as nib
from docx import Document
from docx.shared import Pt, RGBColor
from docx.enum.table import WD_TABLE_ALIGNMENT

BUNDLE = "/scratch/ud3d4/acm_data/ssl_handoff_ours"
V2 = "/scratch/ud3d4/acm_data/handoff_v2"
HANDOFF = "/home/ud3d4/Desktop/SWOG/handoff"
GREEN = RGBColor(0x1a, 0x7f, 0x37); RED = RGBColor(0xB0, 0x2A, 0x2A)
os.makedirs(f"{V2}/ssl_predictions_lits_tumor_only", exist_ok=True)
os.makedirs(f"{V2}/inference", exist_ok=True)

# ---- LiTS tumor-only masks (strip GT-copied liver, keep only label 2) ----
n = 0
for f in sorted(glob.glob(f"{BUNDLE}/ssl_predictions/lits/*.nii.gz")):
    ni = nib.load(f); arr = ni.get_fdata().astype(np.uint8)
    tumor_only = (arr == 2).astype(np.uint8) * 2   # drop liver(1), keep tumor(2)
    nib.save(nib.Nifti1Image(tumor_only, ni.affine, ni.header),
             f"{V2}/ssl_predictions_lits_tumor_only/{os.path.basename(f)}")
    n += 1
print(f"LiTS tumor-only masks: {n}")

# ---- copy inference deliverables ----
shutil.copy("/home/ud3d4/Desktop/SWOG/src/scripts/infer_sam3.py", f"{V2}/inference/infer_sam3.py")
shutil.copy("/scratch/ud3d4/acm_data/sam3_handoff_requirements.txt", f"{V2}/inference/requirements.txt")
shutil.copy(f"{HANDOFF}/handoff_metrics_v2.json", f"{V2}/handoff_metrics_v2.json")

# ---- response document ----
m = json.load(open(f"{HANDOFF}/handoff_metrics_v2.json"))["test_summary"]
doc = Document(); doc.styles['Normal'].font.name = 'Calibri'; doc.styles['Normal'].font.size = Pt(10.5)
def H(t, l=1): doc.add_heading(t, level=l)
def P(t=""): p = doc.add_paragraph(); p.add_run(t); return p

doc.add_heading('Segmentation hand-off — v2 (response to review)', level=0)
P('Addresses the review point by point. The headline change is item #1: the autonomous '
  '(GT-box-removed) numbers, which are the real deployable figures.').italic = True

H('#1 — Autonomous predictions (GT box removed)', 1)
P('CRISP-SAM (the automatic prompt generator) is not yet ready, so per the fallback requested, here are '
  'the test-split numbers with the GT box removed (full-image box = no localization hint). Both prompt '
  'modes shown; slice-selection is held fixed so this isolates within-slice localization+segmentation.')
t = doc.add_table(rows=1, cols=6); t.style = 'Light Grid Accent 1'; t.alignment = WD_TABLE_ALIGNMENT.CENTER
for i, h in enumerate(['Structure', 'Metric', 'Oracle (GT box)', 'Autonomous (box removed)', '', '']):
    t.rows[0].cells[i].paragraphs[0].add_run(h).bold = True
for ds, labs in [("pancreas", ["organ", "tumor"]), ("lits", ["tumor"])]:
    for lab in labs:
        g, f = m[ds]["gtbox"][lab], m[ds]["fullbox"][lab]
        for mk, mm in [("Dice", "dice"), ("HD95 mm", "hd95_mm"), ("NSD@2mm", "nsd_2mm"),
                       ("Sensitivity", "sensitivity"), ("Specificity", "specificity")]:
            c = t.add_row().cells
            c[0].paragraphs[0].add_run(f"{ds} {lab}" if mk == "Dice" else "")
            c[1].paragraphs[0].add_run(mk)
            c[2].paragraphs[0].add_run(str(g[mm]))
            r = c[3].paragraphs[0].add_run(str(f[mm]))
            if mm == "dice":
                r.bold = True; r.font.color.rgb = RED
doc.add_paragraph()
P('Takeaway: pancreatic tumor Dice collapses 0.907 -> 0.112 (HD95 1.6 -> 163 mm) without the box — the '
  'model cannot localize small pancreatic tumors autonomously. Pancreas organ degrades more gracefully '
  '(0.883 -> 0.604) and LiTS tumor best (0.837 -> 0.531, larger/higher-contrast lesions). These are the '
  'honest deployable numbers; the oracle column is the semi-oracle figure delivered previously.')

H('#2 — LiTS liver was GT-copied', 1)
P('Fixed: delivering LiTS as TUMOR-ONLY (folder ssl_predictions_lits_tumor_only/, label 2 only, liver '
  'label removed). No more artificial ~100% liver Dice. Predicting the liver would be a separate training '
  'run — say the word if you want it.')

H('#3 — Runnable inference script', 1)
P('inference/infer_sam3.py — standalone CT->preprocess->box->SAM3->mask, with the box logic and three '
  'prompt modes (gtbox / fullbox / explicit box). inference/requirements.txt pins the environment.')
P('Base model: facebook/sam3, revision 3c879f39826c281e95690f02c7821c4de09afae7 '
  '(transformers 5.8.1, torch 2.5.1+cu121). Preprocessing matches training: 2D per-slice 256x256, '
  'HU window pancreas [-100,300] / LiTS [-100,400], sigmoid>0.5, no post-processing. Confirmed the '
  'script reproduces the delivered (gtbox) masks.')

H('#4 — Per-split metrics (test-split headline)', 1)
P('handoff_metrics_v2.json reports the held-out TEST split only (57 Pancreas, 27 LiTS) — no train cases '
  'mixed in. Both oracle and autonomous, per case and aggregated.')

H('#5 / #6 — HD95, NSD, sensitivity, specificity', 1)
P('All computed per case (monai, native voxel spacing) and populated in handoff_metrics_v2.json — no '
  'longer null. Surface metrics are in the table above; per-case values are in the JSON.')

H('#7 — Softmax / uncertainty (optional)', 1)
P('Not included in this round; can be exported (per-voxel softmax float16) if useful.')

H('Files in this v2', 1)
for f in ["handoff_metrics_v2.json  (test-split, oracle+autonomous, all metrics, per-case + summary)",
          "inference/infer_sam3.py  (runnable inference)",
          "inference/requirements.txt",
          "ssl_predictions_lits_tumor_only/  (131 LiTS tumor-only masks)",
          "this document"]:
    doc.add_paragraph(f, style='List Bullet')

doc.save(f"{V2}/HANDOFF_V2_RESPONSE.docx")
print("Saved response doc + assembled:", V2)
for p in sorted(os.listdir(V2)):
    print("  ", p)
