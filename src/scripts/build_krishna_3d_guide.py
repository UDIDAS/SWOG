#!/usr/bin/env python3
"""Generate the 3D-reconstruction resource guide (.docx, role/artifact-neutral) for the application handoff.
Reads MANIFEST.json for the real case counts. Output: results/3D_Reconstruction_Resources.docx
"""
import json
import os

from docx import Document
from docx.shared import Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH

import glob
PACK = "/scratch/ud3d4/acm_data/krishna_3d_pack"
OUT = "/home/ud3d4/Desktop/SWOG/results/3D_Reconstruction_Resources.docx"
if os.path.exists(f"{PACK}/MANIFEST.json"):
    man = json.load(open(f"{PACK}/MANIFEST.json")); s = man["summary"]
else:                                                 # not finished yet -> count live from disk
    c_ct = len(glob.glob(f"{PACK}/cases_ct/*")); c_pred = len(glob.glob(f"{PACK}/cases_pred/*"))
    c_mesh = len(glob.glob(f"{PACK}/cases_mesh/*"))
    s = {"cases_ct": c_ct, "cases_pred": c_pred, "cases_mesh": c_mesh,
         "total_3d_reconstructions": c_ct + c_pred + c_mesh}
    man = {"summary": s}

d = Document()
d.styles["Normal"].font.name = "Calibri"; d.styles["Normal"].font.size = Pt(10.5)


def h(t, lvl=1):
    p = d.add_heading(t, level=lvl); return p

def body(t):
    return d.add_paragraph(t)

def bullet(t):
    return d.add_paragraph(t, style="List Bullet")


title = d.add_paragraph(); title.alignment = WD_ALIGN_PARAGRAPH.CENTER
r = title.add_run("3D Reconstruction Resources"); r.bold = True; r.font.size = Pt(20)
sub = d.add_paragraph(); sub.alignment = WD_ALIGN_PARAGRAPH.CENTER
sr = sub.add_run("Abdominal CT — organ & tumor 3D surfaces for the imaging application")
sr.italic = True; sr.font.size = Pt(11); sr.font.color.rgb = RGBColor(0x55, 0x55, 0x55)

h("What this pack contains", 1)
body(f"App-ready 3D reconstructions of liver, kidney, pancreas and tumor, built from abdominal CT "
     f"segmentations. Total 3D reconstructions: {s.get('total_3d_reconstructions','—')}.")
tbl = d.add_table(rows=1, cols=3); tbl.style = "Light Grid Accent 1"
hd = tbl.rows[0].cells
hd[0].text, hd[1].text, hd[2].text = "Folder", "Count", "What it gives you"
for folder, cnt, desc in [
    ("cases_ct/", s.get("cases_ct", "—"), "CT volume + ground-truth segmentation + colored meshes + preview. Use these to show a 3D organ model overlaid on the CT."),
    ("cases_pred/", s.get("cases_pred", "—"), "Ground truth AND our model's 3D prediction, both meshed, side-by-side preview. Use to demonstrate model output vs reference."),
    ("cases_mesh/", s.get("cases_mesh", "—"), "Segmentation + meshes only (no CT). A large library of 3D organ reconstructions to populate the app."),
    ("scripts/", "2", "meshify.py and reconstruct_3d.py to regenerate or extend the assets."),
]:
    row = tbl.add_row().cells; row[0].text = str(folder); row[1].text = str(cnt); row[2].text = desc

h("Per-case files", 1)
bullet("CT.nii.gz — the source CT volume (medical NIfTI, carries world affine).")
bullet("seg_gt.nii.gz / seg_pred.nii.gz / seg.nii.gz — the 3D label mask, aligned to the CT affine so it overlays exactly.")
bullet("meshes/gt_organs.glb (and pred_organs.glb) — one colored, app-ready scene with all organs. Drop straight into a web or desktop 3D viewer.")
bullet("meshes/{liver,kidney,pancreas,tumor}.stl — per-organ surface, universal format (Slicer, MeshLab, 3D print, engines).")
bullet("preview.png — a quick render so you can eyeball each case without opening a viewer.")

h("Label & color scheme", 1)
tbl2 = d.add_table(rows=1, cols=3); tbl2.style = "Light Grid Accent 1"
hc = tbl2.rows[0].cells; hc[0].text, hc[1].text, hc[2].text = "Structure", "NIfTI label", "Mesh color"
for name, lab, col in [("Liver", "1", "warm brown"), ("Kidney (L+R merged)", "2 / 13", "purple"),
                       ("Pancreas", "4", "green"), ("Tumor", "14", "red")]:
    rr = tbl2.add_row().cells; rr[0].text, rr[1].text, rr[2].text = name, lab, col

h("Coordinates & units", 1)
bullet("Meshes are in millimetres — voxel spacing is baked in via marching cubes, so sizes are anatomically real.")
bullet("A mesh's local frame has its origin at the volume corner; all organs of one case share that frame, so they assemble correctly together.")
bullet("NIfTI files carry the CT world affine; load CT and seg together and they register voxel-for-voxel.")

h("How to view", 1)
body("Web (simplest) — Google's <model-viewer> renders a GLB with zero 3D code:")
code = d.add_paragraph('<script type="module" src="https://unpkg.com/@google/model-viewer/dist/model-viewer.min.js"></script>\n'
                       '<model-viewer src="cases_ct/FLARE23_0003/meshes/gt_organs.glb"\n'
                       '              camera-controls auto-rotate style="width:600px;height:600px">'
                       '</model-viewer>')
code.style = d.styles["Normal"]
for rn in code.runs:
    rn.font.name = "Consolas"; rn.font.size = Pt(9)
bullet("three.js / Babylon.js / Unity / Unreal — all import GLB natively.")
bullet("Desktop / medical — 3D Slicer or MeshLab open the STL and the NIfTI segmentation directly.")

h("Regenerate or extend", 1)
body("Mesh any segmentation volume (CPU-only):")
c2 = d.add_paragraph("python scripts/meshify.py  SEG.nii.gz  OUT_DIR  --tag gt")
for rn in c2.runs: rn.font.name = "Consolas"; rn.font.size = Pt(9)
body("Reconstruct straight from a raw CT with the trained models (needs one GPU):")
c3 = d.add_paragraph("python scripts/reconstruct_3d.py  CT.nii.gz  OUT_DIR")
for rn in c3.runs: rn.font.name = "Consolas"; rn.font.size = Pt(9)
body("Headroom to scale up: 2,200 full 3D label volumes are available locally (mesh-ready, offline) and "
     "950 of them have retrievable CT — so this library can be expanded well beyond the current set on request.")

d.save(OUT)
print("saved", OUT)
